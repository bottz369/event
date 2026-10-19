"""Deterministic, read-only .xlsx timetable imports; no AI, DB or network calls."""
import io
import math
import re
import unicodedata
import zipfile
from datetime import datetime, time

from models.event_format import new_part, schedule
from models.timetable import TimetableRowDraft

MAX_UPLOAD = 10 * 1024 * 1024
ALIASES = dict(
    artist={'グループ名', 'アーティスト', 'アーティスト名', '出演者'},
    start={'start', '出演開始'}, end={'end', '出演終了'},
    goods_start={'物販開始'}, goods_end={'物販終了'}, place={'物販場所', '場所'},
    duration={'持ち時間', '出演分'}, goods_duration={'物販時間', '物販分'})
BREAKS = {'調整', '転換', '休憩'}


def _key(value):
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', str(value or ''))).casefold()


def _book(data):
    if not data or len(data) > MAX_UPLOAD:
        raise ValueError('エクセルは10MB以下の.xlsxファイルを選択してください。')
    from openpyxl import load_workbook
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
                raise ValueError('エクセルの展開後サイズが上限を超えています。必要なシートだけにしてください。')
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('エクセルを開けませんでした。パスワードなしの.xlsx形式で保存してください。') from exc
    if len(book.worksheets) > 6:
        book.close()
        raise ValueError('エクセルは6シート以内にしてください。')
    return book


def sheet_names(data):
    book = _book(data)
    try:
        return [sheet.title for sheet in book.worksheets]
    finally:
        book.close()


def _minute(value, label):
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        if value.second or value.microsecond:
            raise ValueError(f'{label}：秒を含む時刻には未対応です。分単位にしてください。')
        return value.hour * 60 + value.minute
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if math.isfinite(value) and 0 <= value < 1:
            minutes = value * 1440
            if abs(minutes - round(minutes)) < .00001:
                return round(minutes)
    text = unicodedata.normalize('NFKC', str(value or '')).strip()
    match = re.fullmatch(r'([01]?\d|2[0-3]):([0-5]\d)(?::00)?', text)
    if match:
        return int(match[1]) * 60 + int(match[2])
    raise ValueError(f'{label}：時刻を読み取れません。HH:MM形式で入力し、数式の場合は計算結果を保存してください。')


def _clock(minutes):
    return f'{minutes // 60:02d}:{minutes % 60:02d}'


def _check_minutes(value, expected, label):
    if value is None or value == '':
        return
    try:
        number = float(value)
        if math.isfinite(number) and number == expected:
            return
    except (TypeError, ValueError):
        pass
    raise ValueError(f'{label}：開始・終了時刻と分数が一致しません。')


def parse_rows(values, sheet='シート'):
    """Read labelled columns, preserving source order and explicit goods times."""
    header = None
    entries = []
    for number, cells in enumerate(values, 1):
        cells = list(cells)
        if not any(v is not None and v != '' for v in cells):
            continue
        keys = [_key(v) for v in cells]
        found = {field: [i for i, key in enumerate(keys) if key in aliases]
                 for field, aliases in ALIASES.items()}
        if all(found[k] for k in ('artist', 'start', 'end')):
            if header is not None:
                raise ValueError('1シートに複数の表があります。部ごとにシートを分けて選択してください。')
            if any(len(indices) > 1 for indices in found.values()):
                raise ValueError('同じ意味のヘッダ列が複数あります。列名を確認してください。')
            header = {field: indices[0] for field, indices in found.items() if indices}
            if ('goods_start' in header) != ('goods_end' in header):
                raise ValueError('物販開始・物販終了は両方の列を用意してください。')
            continue
        if header is None:
            continue
        def cell(field):
            index = header.get(field)
            return cells[index] if index is not None and index < len(cells) else None
        if not any(cell(k) is not None and cell(k) != '' for k in header):
            continue
        label = f'{sheet} {number}行目'
        name = str(cell('artist') or '').strip()
        if not name:
            raise ValueError(f'{label}：アーティスト名が空欄です。')
        start, end = _minute(cell('start'), label+' 開始'), _minute(cell('end'), label+' 終了')
        if end <= start:
            raise ValueError(f'{label}：終了は開始より後にしてください。日付をまたぐ時間帯には未対応です。')
        _check_minutes(cell('duration'), end-start, label+' 持ち時間')
        if entries and start < entries[-1]['end']:
            raise ValueError(f'{label}：出演時間が前の行と重なっています。')
        gs, ge = cell('goods_start'), cell('goods_end')
        has_gs, has_ge = gs not in (None, ''), ge not in (None, '')
        if has_gs != has_ge:
            raise ValueError(f'{label}：物販開始・終了は両方入力してください。')
        goods_start = _minute(gs, label+' 物販開始') if has_gs else None
        goods_end = _minute(ge, label+' 物販終了') if has_ge else None
        if has_gs and goods_end < goods_start:
            raise ValueError(f'{label}：物販終了が開始より前です。')
        goods_duration = goods_end-goods_start if has_gs else 0
        _check_minutes(cell('goods_duration'), goods_duration, label+' 物販時間')
        is_break = name in BREAKS
        if is_break and (has_gs or str(cell('place') or '').strip()):
            raise ValueError(f'{label}：調整行に物販情報があります。出演者名を確認してください。')
        entries.append(dict(name=name, start=start, end=end, goods_start=goods_start,
                            goods_duration=goods_duration, place=str(cell('place') or '').strip(), is_break=is_break))
    if header is None:
        raise ValueError('ヘッダ行にSTART・END・グループ名（または出演開始・出演終了・アーティスト）を用意してください。')
    artists = [entry for entry in entries if not entry['is_break']]
    if not artists:
        raise ValueError('出演者を読み取れませんでした。')
    if len(artists) > 500:
        raise ValueError('出演者は500行以内にしてください。')
    part = new_part(0, '', _clock(artists[0]['start']))
    warnings = []
    for index, entry in enumerate(artists):
        gap = artists[index+1]['start']-entry['end'] if index+1 < len(artists) else 0
        row = TimetableRowDraft(artist_name=entry['name'], duration=entry['end']-entry['start'],
                                adjustment=gap, goods_start_time=_clock(entry['goods_start']) if entry['goods_start'] is not None else '',
                                goods_duration=entry['goods_duration'], place=entry['place'])
        part['rows'].append(row.to_legacy_dict())
        if entry['goods_start'] is not None and entry['goods_start'] < entry['start']:
            warnings.append(f"{entry['name']}：物販が出演前の時刻です。エクセルの時刻をそのまま取り込みます。")
    calculated, _ = schedule(part)
    if [(r['start'], r['end']) for r in calculated] != [(_clock(e['start']), _clock(e['end'])) for e in artists]:
        raise ValueError('アプリの計算結果とエクセルの出演時刻が一致しません。')
    count = sum(entry['is_break'] for entry in entries)
    if count:
        warnings.insert(0, f'調整・転換・休憩の{count}行は出演者に含めず、出演者間の空き時間を転換分に反映します。')
    return dict(parts=[part], warnings=warnings, sheet=sheet, unmatched=[])


def read_excel(data, sheet):
    book = _book(data)
    try:
        if sheet not in book.sheetnames:
            raise ValueError('選択したシートがありません。')
        selected = book[sheet]
        if selected.max_row > 1000 or selected.max_column > 40:
            raise ValueError('1シート1000行・40列以内にしてください。')
        return parse_rows(selected.iter_rows(values_only=True), sheet)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('シートを読み取れませんでした。数式の計算結果を保存した.xlsxでお試しください。') from exc
    finally:
        book.close()
